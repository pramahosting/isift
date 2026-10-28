import base64, io, json
import pytest
from app.services import docconvert as dc



def _png_url():
    import PIL.Image as I
    b = io.BytesIO(); I.new("RGB", (80, 60), (220, 70, 70)).save(b, "PNG")
    return "data:image/png;base64," + base64.b64encode(b.getvalue()).decode()


HTML = '<h1>Review <span style="color:#c0392b">2026</span></h1><p>Plain <b>bold</b> <i>italic</i> <a href="https://isift.au">link</a></p><ul><li>One</li><li>Two</li></ul><table border="1"><tbody><tr><td>A</td><td>B</td></tr><tr><td>1</td><td>2</td></tr></tbody></table><script>alert(1)</script>'


def test_wordperfect_to_pdf_and_back_keeps_headings_bold_colour_lists_and_tables():
    pdf = dc.wp_to_pdf(HTML, "Review")
    assert pdf[:5] == b"%PDF-"
    back = dc.pdf_to_wp(pdf)["html"]
    import re as _re
    assert "<h1" in back and "Review" in back and _re.search(r"<b>[^<]*(<span[^>]*>)?bold", back) and _re.search(r"<i>[^<]*(<span[^>]*>)?italic", back)
    assert "#bf382a" in back or "#c0392b" in back                          # the red keeps its colour
    assert "<ul>" in back and "One" in back and "Two" in back and "<table" in back and _re.search(r"<td[^>]*>(<span[^>]*>)?A\b", back) and _re.search(r"<td[^>]*>(<span[^>]*>)?B\b", back)
    assert "alert" not in back                                              # scripts never reach the PDF


def test_wordperfect_to_docx_is_real_word_content():
    from docx import Document
    d = Document(io.BytesIO(dc.wp_to_docx(HTML + f'<p><img src="{_png_url()}" width="80" height="60"></p>', "Review")))
    styles = [p.style.name for p in d.paragraphs]
    assert "Heading 1" in styles and "List Bullet" in styles and len(d.tables) == 1 and len(d.inline_shapes) == 1
    runs = [r for p in d.paragraphs for r in p.runs]
    assert any(r.bold for r in runs) and any(r.italic for r in runs)


def test_spreadtable_to_xlsx_keeps_numbers_text_and_formulas():
    import openpyxl
    book = json.dumps({"sheets": [{"name": "Budget", "cells": {"A1": "Item", "B1": "Cost", "B2": "1200", "B3": "450.5", "B4": "=SUM(B2:B3)", "C1": "007", "D1": "=CONCAT(A1,B1)"}}, {"name": "Budget", "cells": {"A1": "x"}}]})
    wb = openpyxl.load_workbook(io.BytesIO(dc.sheet_to_xlsx(book, "B")))
    assert wb.sheetnames == ["Budget", "Budget2"]
    ws = wb["Budget"]
    assert ws["B2"].value == 1200 and ws["B3"].value == 450.5 and ws["B4"].value == "=SUM(B2:B3)" and ws["C1"].value == "007" and "CONCAT" in ws["D1"].value


def test_spreadtable_to_pdf_shows_worked_out_values():
    from pypdf import PdfReader
    book = json.dumps({"sheets": [{"name": "S", "cells": {"A1": "Cost", "B1": "=SUM(1,2)"}}]})
    t = PdfReader(io.BytesIO(dc.sheet_to_pdf(book, [{"B1": "3"}], "S"))).pages[0].extract_text()
    assert "Cost" in t and "3" in t and "SUM" not in t


DECK = json.dumps({"theme": "dark", "slides": [
    {"layout": "title", "title": "iSift", "body": ["Meet, minute, sign"], "notes": "say hello"},
    {"layout": "content", "title": "Why", "body": ["One app", "90 minutes free"], "shapes": [{"id": "s1", "lib": "x", "x": 600, "y": 300, "w": 200}]},
    {"layout": "two", "title": "Two", "body": ["a"], "body2": ["c"], "bg": "#fff7ed"}]})


def test_presentation_to_pdf_one_page_per_slide_with_text_and_shapes():
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(dc.deck_to_pdf(DECK, {"s1": {"png": _png_url(), "w": 80, "h": 60}}, "D")))
    assert len(r.pages) == 3 and "Why" in r.pages[1].extract_text() and "90 minutes free" in r.pages[1].extract_text()
    assert float(r.pages[0].mediabox.width) > float(r.pages[0].mediabox.height)


def test_presentation_to_pptx_has_text_shapes_and_notes():
    from pptx import Presentation
    p = Presentation(io.BytesIO(dc.deck_to_pptx(DECK, {"s1": {"png": _png_url(), "w": 80, "h": 60}}, "D")))
    assert len(p.slides) == 3 and p.slides[0].notes_slide.notes_text_frame.text == "say hello"
    texts = [sh.text_frame.text for sh in p.slides[1].shapes if sh.has_text_frame]
    assert "Why" in texts and any("One app" in t for t in texts) and any(sh.shape_type == 13 for sh in p.slides[1].shapes)      # a picture


def test_pictures_only_from_inside_the_file_and_no_web_fetch():
    pdf = dc.wp_to_pdf('<img src="https://example.com/x.png"><img src="file:///etc/passwd"><p>ok</p>', "T")
    assert pdf[:5] == b"%PDF-"


def test_pdf_with_no_text_becomes_pictures_and_not_a_pdf_is_refused():
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument.new(); pdf.new_page(200, 200); b = io.BytesIO(); pdf.save(b)
    r = dc.pdf_to_wp(b.getvalue())
    assert "<img" in r["html"] and r["notes"]
    with pytest.raises(dc.ConvertError): dc.pdf_to_wp(b"hello")


import asyncio
from starlette.testclient import TestClient


@pytest.fixture()
def client():
    from app.main import app
    from app.core.db import engine
    from app.models import Base
    async def mk():
        async with engine.begin() as c: await c.run_sync(Base.metadata.create_all)
    asyncio.run(mk())
    with TestClient(app) as c: yield c


def _auth(c):
    r = c.post("/api/v1/auth/local-login", json={"username": "admin", "password": "Meeting$1"}); assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["access_token"]}


def test_convert_route_needs_sign_in_and_returns_the_file(client):
    assert client.post("/api/v1/docs/convert", json={"kind": "wp", "to": "pdf", "content": "<p>x</p>"}).status_code in (401, 403)
    h = _auth(client)
    r = client.post("/api/v1/docs/convert", json={"kind": "wp", "to": "docx", "title": "My: doc/1", "content": HTML}, headers=h)
    assert r.status_code == 200, r.text
    j = r.json(); assert j["filename"] == "My doc1.docx" and base64.b64decode(j["data"])[:2] == b"PK" and "wordprocessingml" in j["mime"]
    assert client.post("/api/v1/docs/convert", json={"kind": "wp", "to": "xlsx", "content": "x"}, headers=h).status_code == 422
    p = client.post("/api/v1/docs/convert", json={"kind": "wp", "to": "pdf", "content": HTML}, headers=h).json()
    back = client.post("/api/v1/docs/pdf-to-wp", json={"name": "a.pdf", "data": p["data"]}, headers=h)
    assert back.status_code == 200 and "<h1" in back.json()["html"]
    assert client.post("/api/v1/docs/pdf-to-wp", json={"data": base64.b64encode(b"not a pdf").decode()}, headers=h).status_code == 422


def _office_files():
    import io
    from docx import Document
    from docx.shared import Pt, RGBColor
    from openpyxl import Workbook
    from pptx import Presentation
    d = Document(); p = d.add_paragraph(); r = p.add_run("Red words"); r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B); r.font.size = Pt(16); r.bold = True
    d.add_paragraph("one", style="List Bullet")
    b = io.BytesIO(); d.save(b); docx = b.getvalue()
    wb = Workbook(); ws = wb.active; ws.title = "Budget"; ws["A1"] = "Item"; ws["B1"] = 5; ws["B2"] = 7.5; ws["B3"] = "=SUM(B1:B2)"
    b = io.BytesIO(); wb.save(b); xlsx = b.getvalue()
    pr = Presentation(); s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = "Hello"; s.placeholders[1].text = "one\ntwo"; s.notes_slide.notes_text_frame.text = "say hi"
    b = io.BytesIO(); pr.save(b); return docx, xlsx, b.getvalue()


def test_microsoft_files_open_and_keep_their_look():
    from app.services import docconvert as dc, htmlsafe
    docx, xlsx, pptx = _office_files()
    h = htmlsafe.clean(dc.docx_to_wp(docx)["html"])
    assert "color:#c0392b" in h and "font-size:16pt" in h and "<b>" in h and "<ul>" in h       # colour, size, bold and the list survive being saved
    book = json.loads(dc.xlsx_to_book(xlsx)["content"])
    assert book["sheets"][0]["name"] == "Budget" and book["sheets"][0]["cells"]["B3"] == "=SUM(B1:B2)" and book["sheets"][0]["cells"]["B2"] == "7.5"
    deck = json.loads(dc.pptx_to_deck(pptx)["content"])
    assert deck["slides"][0]["title"] == "Hello" and deck["slides"][0]["body"] == ["one", "two"] and deck["slides"][0]["notes"] == "say hi"
    with pytest.raises(dc.ConvertError): dc.docx_to_wp(b"not a word file")


def test_import_route_and_roundtrip_to_word(client):
    docx, xlsx, pptx = _office_files()
    c = client; h = _auth(c)
    for name, raw, kind in (("a.docx", docx, "wp"), ("b.xlsx", xlsx, "sheet"), ("c.pptx", pptx, "deck")):
        r = c.post("/api/v1/docs/import", json={"name": name, "data": base64.b64encode(raw).decode()}, headers=h); assert r.status_code == 200 and r.json()["kind"] == kind, r.text
    assert c.post("/api/v1/docs/import", json={"name": "x.doc", "data": base64.b64encode(docx).decode()}, headers=h).status_code == 422
    html = c.post("/api/v1/docs/import", json={"name": "a.docx", "data": base64.b64encode(docx).decode()}, headers=h).json()["content"]
    out = c.post("/api/v1/docs/convert", json={"kind": "wp", "to": "docx", "title": "t", "content": html}, headers=h).json()
    from docx import Document
    run = [r for p in Document(io.BytesIO(base64.b64decode(out["data"]))).paragraphs for r in p.runs if r.text == "Red words"][0]
    assert str(run.font.color.rgb) == "C0392B" and run.bold and run.font.size.pt == 16           # opened, edited and saved again: same colour, size, bold


def test_pdf_to_wordperfect_keeps_colours_and_shading():
    from weasyprint import HTML
    from app.services import docconvert as dc, htmlsafe
    pdf = HTML(string='<h1 style="color:#c0392b;font:bold 26pt Arial">Title</h1><div style="background:#1a56db;color:#fff;padding:8px;font:bold 14pt Arial">Bar</div><p>Text <span style="color:#188038;font-size:18pt">green</span></p>').write_pdf()
    h = htmlsafe.clean(dc.pdf_to_wp(pdf)["html"])
    assert "color:#bf382a" in h or "color:#c0392b" in h
    assert "background-color:#1a56db" in h and "color:#ffffff" in h and "color:#188038" in h and "font-size:18pt" in h
